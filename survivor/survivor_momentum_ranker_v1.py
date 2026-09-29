"""
Survivor Momentum Ranker v1

Narrow 9-symbol monitor for the retained M4/P7 strategy.

Requires the existing file:
    C:\Users\wmoon\Documents\MATLAB\nefty_100_momentum_ranker_v10_hil_semantics.py

The existing v10 file is reused only for:
- TradeStation OAuth/session handling
- daily-bar retrieval
- verified TS-close x Yahoo adjustment-factor total-return construction
- rolling current momentum calculations
- GitHub upload helper

Survivor replaces the broad NEFTy universe and all NEFTy selection logic.

Risky universe:
    SPY, QQQ, TLT, GLD, HYG, XLE, TIP, UUP
Reserve:
    SHY

Official monthly authority:
- 12-1 momentum
- top 3 positive risky assets
- fixed 3 slots; vacancies go to SHY
- canonical P7 tilt only for selected SPY/QQQ
- current daily horizons are surveillance only and do not alter monthly targets

No orders are submitted.
"""

# BUILD: SURVIVOR_MOMENTUM_RANKER_V1_20260929

import argparse
import csv
import importlib.util
import shutil
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

BASE_FILE = Path(
    r"C:\Users\wmoon\Documents\MATLAB\nefty_100_momentum_ranker_v10_hil_semantics.py"
)
RUNTIME = Path(r"D:\Algo_Trading\runtime_state")
DEFAULT_CSV = RUNTIME / "Survivor_LATEST.csv"
DEFAULT_XLSX = RUNTIME / "Survivor_LATEST.xlsx"

RISKY = ["SPY", "QQQ", "TLT", "GLD", "HYG", "XLE", "TIP", "UUP"]
SHY = "SHY"
ALL = RISKY + [SHY]
TOP_N = 3
P7_MAX_TILT = 0.50

AXIS = {
    "SPY": "U.S. Equity",
    "QQQ": "U.S. Equity",
    "TLT": "Nominal Duration",
    "GLD": "Precious Metals",
    "HYG": "Credit / Risk Appetite",
    "XLE": "Energy",
    "TIP": "Inflation-Linked Duration",
    "UUP": "USD / Currency",
    "SHY": "Reserve",
}

EXPRESSION = {
    "SPY": "S&P 500 broad equity",
    "QQQ": "Nasdaq-100 growth / technology",
    "TLT": "20+ Year U.S. Treasury",
    "GLD": "Gold bullion",
    "HYG": "High-yield corporate credit",
    "XLE": "U.S. energy equities",
    "TIP": "U.S. inflation-protected Treasuries",
    "UUP": "U.S. dollar bullish basket",
    "SHY": "1-3 Year U.S. Treasury reserve",
}

HEADERS = [
    "Rank",
    "ETF",
    "Role",
    "Axis",
    "Expression",
    "18m",
    "12m",
    "12-1",
    "6m",
    "3m",
    "1m",
    "1d",
    "12-1 Positive",
    "Official Signal Month",
    "Official 12-1",
    "Official Rank",
    "Official Selected",
    "Official Target Weight",
    "Official Status",
]


def load_base():
    if not BASE_FILE.exists():
        raise FileNotFoundError(
            f"Required base data module not found: {BASE_FILE}\n"
            "Keep the existing v10 ranker in C:\\Users\\wmoon\\Documents\\MATLAB."
        )

    spec = importlib.util.spec_from_file_location("nefty_v10_base", BASE_FILE)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)

    # Restrict all reused data functions to Survivor only.
    mod.ETF_UNIVERSE = list(ALL)
    return mod


def build_month_end_levels(total_return_series):
    levels = {sym: {} for sym in ALL}

    for sym in ALL:
        for date, value in total_return_series[sym]:
            # Series is date-sorted. Reassignment leaves the last observed
            # trading-session value for each calendar month.
            levels[sym][date[:7]] = value

    common_months = sorted(
        set.intersection(*(set(levels[sym]) for sym in RISKY))
    )
    return common_months, levels


def official_signal(common_months, levels):
    current_ym = datetime.now(
        ZoneInfo("America/Los_Angeles")
    ).strftime("%Y-%m")

    completed = [ym for ym in common_months if ym < current_ym]
    if not completed:
        raise RuntimeError("No completed calendar month available.")

    signal_month = completed[-1]
    i = common_months.index(signal_month)

    if i < 12:
        raise RuntimeError(
            "Insufficient month-end history for official Survivor 12-1."
        )

    # Canonical M4 12-1:
    # numerator = month-end immediately before the signal month
    # denominator = month-end 12 months before the signal month
    numerator_month = common_months[i - 1]
    denominator_month = common_months[i - 12]

    scores = {}
    for sym in RISKY:
        p_num = levels[sym].get(numerator_month)
        p_den = levels[sym].get(denominator_month)
        if p_num is None or p_den is None or p_den <= 0:
            raise RuntimeError(
                f"Missing official Survivor signal history for {sym}"
            )
        scores[sym] = p_num / p_den - 1.0

    return signal_month, numerator_month, denominator_month, scores


def survivor_targets(scores):
    positive = {sym: value for sym, value in scores.items() if value > 0}

    selected = sorted(
        positive,
        key=lambda sym: (-positive[sym], sym),
    )[:TOP_N]

    official_rank = {
        sym: rank
        for rank, sym in enumerate(
            sorted(RISKY, key=lambda s: (-scores[s], s)),
            start=1,
        )
    }

    if not selected:
        return selected, official_rank, {SHY: 1.0}

    max_positive = max(positive.values())
    base = 1.0 / len(selected)
    raw = {}

    for sym in selected:
        if sym in ("SPY", "QQQ") and max_positive > 0:
            tilt_multiplier = (
                1.0
                + (scores[sym] / max_positive) * P7_MAX_TILT
            )
            raw[sym] = base * tilt_multiplier
        else:
            raw[sym] = base

    raw_total = sum(raw.values())
    relative = {
        sym: weight / raw_total
        for sym, weight in raw.items()
    }

    # Fixed 3-slot architecture:
    # P7 changes relative weights only inside occupied risky slots.
    # Vacant slots remain SHY and are not redistributed.
    risky_budget = len(selected) / TOP_N

    targets = {
        sym: relative[sym] * risky_budget
        for sym in selected
    }

    if risky_budget < 1.0:
        targets[SHY] = 1.0 - risky_budget

    if abs(sum(targets.values()) - 1.0) > 1e-10:
        raise RuntimeError(
            "Survivor target weights do not sum to 100%."
        )

    return selected, official_rank, targets


def build_current_matrix(base, total_return_series):
    matrix = base.compute_matrix(total_return_series)
    daily = base.compute_daily_changes(total_return_series)

    for sym in ALL:
        matrix[sym]["1d"] = daily.get(sym)

    return matrix


def current_rank_map(matrix):
    valid = [
        sym for sym in RISKY
        if matrix[sym].get("12-1") is not None
    ]

    ordered = sorted(
        valid,
        key=lambda sym: (-matrix[sym]["12-1"], sym),
    )

    return {
        sym: rank
        for rank, sym in enumerate(ordered, start=1)
    }


def build_rows(
    matrix,
    rank_map,
    signal_month,
    official_scores,
    official_rank,
    selected,
    targets,
):
    rows = []
    selected_set = set(selected)

    for sym in ALL:
        if sym in RISKY:
            current_12_1 = matrix[sym].get("12-1")
            off_selected = sym in selected_set
            off_score = official_scores[sym]

            if off_selected:
                off_status = "SELECTED"
            elif off_score > 0:
                off_status = "POSITIVE_NOT_SELECTED"
            else:
                off_status = "NON_POSITIVE"

            row = {
                "Rank": rank_map.get(sym),
                "ETF": sym,
                "Role": "Risky",
                "Axis": AXIS[sym],
                "Expression": EXPRESSION[sym],
                "18m": matrix[sym].get("18m"),
                "12m": matrix[sym].get("12m"),
                "12-1": current_12_1,
                "6m": matrix[sym].get("6m"),
                "3m": matrix[sym].get("3m"),
                "1m": matrix[sym].get("1m"),
                "1d": matrix[sym].get("1d"),
                "12-1 Positive": (
                    current_12_1 is not None
                    and current_12_1 > 0
                ),
                "Official Signal Month": signal_month,
                "Official 12-1": off_score,
                "Official Rank": official_rank[sym],
                "Official Selected": off_selected,
                "Official Target Weight": targets.get(sym, 0.0),
                "Official Status": off_status,
            }
        else:
            row = {
                "Rank": None,
                "ETF": SHY,
                "Role": "Reserve",
                "Axis": AXIS[SHY],
                "Expression": EXPRESSION[SHY],
                "18m": matrix[SHY].get("18m"),
                "12m": matrix[SHY].get("12m"),
                "12-1": matrix[SHY].get("12-1"),
                "6m": matrix[SHY].get("6m"),
                "3m": matrix[SHY].get("3m"),
                "1m": matrix[SHY].get("1m"),
                "1d": matrix[SHY].get("1d"),
                "12-1 Positive": None,
                "Official Signal Month": signal_month,
                "Official 12-1": None,
                "Official Rank": None,
                "Official Selected": False,
                "Official Target Weight": targets.get(SHY, 0.0),
                "Official Status": (
                    "RESERVE"
                    if targets.get(SHY, 0.0) > 0
                    else "RESERVE_UNUSED"
                ),
            }

        rows.append(row)

    return rows


def save_csv(rows, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=HEADERS,
        )
        writer.writeheader()
        writer.writerows(rows)


def save_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Survivor"

    ws.append(HEADERS)

    for row in rows:
        ws.append(
            [row[header] for header in HEADERS]
        )

    header_fill = PatternFill(
        "solid",
        fgColor="1F2937",
    )
    selected_fill = PatternFill(
        "solid",
        fgColor="D9EAD3",
    )
    reserve_fill = PatternFill(
        "solid",
        fgColor="FFF2CC",
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )

    pct_cols = [
        6, 7, 8, 9, 10, 11, 12, 15, 18
    ]

    for row_idx in range(2, ws.max_row + 1):
        is_selected = bool(
            ws.cell(row_idx, 17).value
        )
        is_shy = (
            ws.cell(row_idx, 2).value == SHY
        )

        fill = None
        if is_selected:
            fill = selected_fill
        elif is_shy:
            fill = reserve_fill

        if fill is not None:
            for col_idx in range(
                1,
                ws.max_column + 1,
            ):
                ws.cell(
                    row_idx,
                    col_idx,
                ).fill = fill

        for col_idx in pct_cols:
            ws.cell(
                row_idx,
                col_idx,
            ).number_format = (
                '0.0%;[Red]-0.0%'
            )

    widths = [
        7, 8, 10, 26, 34,
        11, 11, 11, 11, 11, 11, 11,
        14, 20, 14, 14, 16, 20, 24,
    ]

    for idx, width in enumerate(
        widths,
        start=1,
    ):
        ws.column_dimensions[
            get_column_letter(idx)
        ].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)


def print_summary(
    selected,
    targets,
    signal_month,
    numerator_month,
    denominator_month,
):
    print()
    print("=" * 82)
    print(
        "SURVIVOR - MONTHLY AUTHORITY "
        "+ DAILY SURVEILLANCE"
    )
    print(
        f"Official signal month: "
        f"{signal_month}; "
        f"12-1 = {numerator_month}/"
        f"{denominator_month} - 1"
    )
    print("Official targets:")

    for sym, weight in sorted(
        targets.items(),
        key=lambda item: (
            -item[1],
            item[0],
        ),
    ):
        print(
            f"  {sym:<4} "
            f"{weight:>7.1%}"
        )

    print(
        "Selected: "
        + (
            ", ".join(selected)
            if selected
            else "None"
        )
    )

    print(
        "Daily shorter-horizon readings are "
        "surveillance only. No 1m gate."
    )
    print("=" * 82)


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Survivor 9-symbol "
            "momentum monitor"
        )
    )

    parser.add_argument(
        "--csv",
        default=str(DEFAULT_CSV),
    )
    parser.add_argument(
        "--xlsx",
        default=str(DEFAULT_XLSX),
    )
    parser.add_argument(
        "--no-github",
        action="store_true",
    )

    args = parser.parse_args()

    stamp = datetime.now(
        ZoneInfo("America/Los_Angeles")
    ).strftime(
        "%Y-%m-%d_%H%M_PT"
    )

    archive_dir = (
        Path(args.csv).parent
        / "archive"
    )

    archive_csv = (
        archive_dir
        / f"Survivor_{stamp}.csv"
    )
    archive_xlsx = (
        archive_dir
        / f"Survivor_{stamp}.xlsx"
    )

    base = load_base()

    print(
        "[1/6] Authenticating "
        "to TradeStation..."
    )
    session = base.load_session()

    print(
        "[2/6] Fetching daily bars "
        "for 9 Survivor symbols..."
    )
    bars = base.fetch_bars(session)

    missing = [
        sym for sym in ALL
        if sym not in bars
    ]
    if missing:
        raise RuntimeError(
            f"Missing Survivor TS bars: "
            f"{missing}"
        )

    print(
        "[3/6] Building verified "
        "total-return series..."
    )
    adjustment_data, failures = (
        base.fetch_adjustment_data(bars)
    )

    if failures:
        raise RuntimeError(
            "Corporate-action verification "
            f"failures: {failures}"
        )

    total_return_series = (
        base.build_total_return_series(
            bars,
            adjustment_data,
        )
    )

    missing_total_return = [
        sym for sym in ALL
        if sym not in total_return_series
    ]
    if missing_total_return:
        raise RuntimeError(
            "Missing Survivor total-return "
            f"series: {missing_total_return}"
        )

    print(
        "[4/6] Computing current "
        "surveillance and official "
        "monthly target..."
    )

    matrix = build_current_matrix(
        base,
        total_return_series,
    )
    rank_map = current_rank_map(
        matrix
    )

    common_months, levels = (
        build_month_end_levels(
            total_return_series
        )
    )

    (
        signal_month,
        numerator_month,
        denominator_month,
        official_scores,
    ) = official_signal(
        common_months,
        levels,
    )

    (
        selected,
        official_rank,
        targets,
    ) = survivor_targets(
        official_scores
    )

    rows = build_rows(
        matrix,
        rank_map,
        signal_month,
        official_scores,
        official_rank,
        selected,
        targets,
    )

    if (
        len(rows) != 9
        or {
            row["ETF"]
            for row in rows
        } != set(ALL)
    ):
        raise RuntimeError(
            "Survivor output failed "
            "exact 9-symbol validation."
        )

    print(
        "[5/6] Writing local stable "
        "+ immutable outputs..."
    )

    save_csv(
        rows,
        args.csv,
    )
    save_xlsx(
        rows,
        args.xlsx,
    )

    archive_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    shutil.copy2(
        args.csv,
        archive_csv,
    )
    shutil.copy2(
        args.xlsx,
        archive_xlsx,
    )

    print(f"[SAVED] {args.csv}")
    print(f"[SAVED] {args.xlsx}")
    print(f"[SAVED] {archive_csv}")
    print(f"[SAVED] {archive_xlsx}")

    if not args.no_github:
        print(
            "[6/6] Publishing Survivor "
            "files to GitHub..."
        )

        token = (
            base._load_github_token()
        )

        base.github_upload(
            args.csv,
            "Survivor_LATEST.csv",
            token,
        )
        base.github_upload(
            args.xlsx,
            "Survivor_LATEST.xlsx",
            token,
        )
        base.github_upload(
            str(archive_csv),
            f"archive/{archive_csv.name}",
            token,
        )
        base.github_upload(
            str(archive_xlsx),
            f"archive/{archive_xlsx.name}",
            token,
        )
    else:
        print(
            "[6/6] GitHub publish "
            "suppressed."
        )

    print_summary(
        selected,
        targets,
        signal_month,
        numerator_month,
        denominator_month,
    )


if __name__ == "__main__":
    main()
